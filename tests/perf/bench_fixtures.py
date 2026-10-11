"""Central generators for benchmark fixtures: sheets, videos, manifests.

Every benchmark (``scan_bench``, ``clip_bench``, ``tests/ui/ui_bench.py``)
and the profile skill's recipes build their inputs here, so the geometry a
page is measured against is written once. The workbook layout mirrors
``tests/ui/_ui_fixtures._make_workbook`` — ``ID`` at F2 with participant
columns to its right *on row 2*, ``Observation``/``Category`` on row 5, data
from row 6 — because ``spreadsheet.build_sheet_context`` scans for exactly
those labels and a drifted layout yields a silently small grid, not an error.
The manifest sections copy the shapes ``_ui_fixtures`` seeds, scaled up.

Pure: no clipgen imports, so the ``tests/ui`` harness and ``tests/`` unit
tests can load it with ``importlib`` without the source tree on ``sys.path``.
"""

from __future__ import annotations

import json
import shutil
import subprocess
from collections.abc import Callable
from pathlib import Path
from typing import Any

SEVERITIES = ("Critical", "Serious", "Moderate", "Minor")
HEADERS = ("Count", "Reported", "Severity", "Category", "Observation", "Summary")


def write_sheet(
    path: Path,
    *,
    study: str,
    rows: int,
    participants: int,
    cell: Callable[[int, int], str] | None = None,
) -> Path:
    """Write an ``.xlsx`` with *rows* observations across *participants*.

    *cell(row_index, participant_index)* returns the timestamp text for that
    cell; the default fills every third participant column with ``0:01-0:04``
    (the gridbench recipe). Severities cycle so ``renderGrid`` paints its
    ``.sev-*`` classes — an empty Severity column under-measures it.
    """
    import openpyxl

    def default_cell(_row: int, participant: int) -> str:
        return "0:01-0:04" if participant % 3 == 0 else ""

    fill = cell or default_cell
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Observations"
    ws["A1"] = study
    ws["F2"] = "ID"
    for i in range(participants):
        ws.cell(2, 7 + i, f"P{i + 1:02d}")
    for col, header in enumerate(HEADERS, 1):
        ws.cell(5, col, header)
    for r in range(rows):
        ws.cell(6 + r, 3, SEVERITIES[r % len(SEVERITIES)])
        ws.cell(6 + r, 4, "Onboarding")
        ws.cell(6 + r, 5, f"Observation {r}")
        for i in range(participants):
            ws.cell(6 + r, 7 + i, fill(r, i))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def make_testsrc_video(
    path: Path,
    *,
    duration: int,
    size: str = "1280x720",
    rate: int = 30,
    audio: bool = False,
    fragmented: bool = False,
) -> Path:
    """Encode a deterministic testsrc clip (constant motion defeats phash-skip).

    *audio* adds a sine track (Whisper and the clip pipeline need one);
    *fragmented* writes the OBS-style MP4 browsers cannot seek.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    cmd = [
        "ffmpeg",
        "-y",
        "-loglevel",
        "error",
        "-f",
        "lavfi",
        "-i",
        f"testsrc=duration={duration}:size={size}:rate={rate}",
    ]
    if audio:
        cmd += ["-f", "lavfi", "-i", f"sine=frequency=220:duration={duration}"]
    cmd += ["-pix_fmt", "yuv420p", "-c:v", "libx264", "-g", str(rate)]
    if audio:
        cmd += ["-c:a", "aac", "-shortest"]
    if fragmented:
        cmd += ["-movflags", "frag_keyframe+empty_moov+default_base_moof"]
    cmd.append(str(path))
    subprocess.run(cmd, check=True)
    return path


# Two speakers so the speaker pass has turns to cluster.
SPEECH_LINES = (
    (
        "Samantha",
        "Okay, so I am on the onboarding screen and I cannot find the continue button.",
    ),
    ("Daniel", "Take your time. What would you expect to happen next?"),
    (
        "Samantha",
        "I guess I would tap the arrow in the corner, but it looks disabled to me.",
    ),
    ("Daniel", "Why does it look disabled? Tell me what you are seeing."),
    (
        "Samantha",
        "It is grey and very small, and nothing happens when I hover over it.",
    ),
    ("Daniel", "Thanks. Let us move on to the settings page for a moment."),
)


def _say_line(voice: str, text: str, out: Path) -> None:
    """Render *text* with macOS ``say``; the default voice when *voice* is missing."""
    proc = subprocess.run(
        ["say", "-v", voice, "-o", str(out), text], capture_output=True, check=False
    )
    if proc.returncode or not out.is_file():
        subprocess.run(["say", "-o", str(out), text], check=True)


def make_speech_audio(path: Path) -> Path:
    """A ~40 s two-voice dialogue (macOS ``say``) with pauses between turns.

    Real speech, so VAD, Whisper and the speaker pass all do their actual
    work — a sine track is all silence to VAD and measures nothing.
    """
    if shutil.which("say") is None:
        raise RuntimeError("speech fixtures need macOS 'say'; pass --source instead")
    path.parent.mkdir(parents=True, exist_ok=True)
    parts: list[Path] = []
    for i, (voice, text) in enumerate(SPEECH_LINES):
        part = path.parent / f"_line{i}.aiff"
        _say_line(voice, text, part)
        parts.append(part)
    inputs: list[str] = []
    for part in parts:
        inputs += ["-i", str(part)]
    pads = "".join(f"[{i}:a]apad=pad_dur=0.8[a{i}];" for i in range(len(parts)))
    joined = "".join(f"[a{i}]" for i in range(len(parts)))
    graph = f"{pads}{joined}concat=n={len(parts)}:v=0:a=1[out]"
    subprocess.run(
        ["ffmpeg", "-y", "-loglevel", "error", *inputs]
        + ["-filter_complex", graph, "-map", "[out]", "-ar", "48000", str(path)],
        check=True,
    )
    for part in parts:
        part.unlink()
    return path


def make_long_video(
    path: Path,
    *,
    duration: int,
    size: str = "1920x1080",
    rate: int = 30,
    gop: int = 250,
    audio: str = "sine",
    content: str = "screen",
) -> Path:
    """An hours-long OBS-like recording, built in seconds by looping a short seed.

    The seed holds a whole number of fixed-length GOPs (``-g``/``-keyint_min``
    with scene-cut off, as OBS writes them) and is stream-copied with
    ``-stream_loop`` up to *duration*, so a 3 h fixture costs one ~50 s
    encode. *content* is ``screen`` (a still frame changing every 5 s, like
    a screen recording, so phash-skip behaves as on real footage) or
    ``testsrc`` (constant motion, the worst case). *audio* is ``sine``,
    ``speech`` (:func:`make_speech_audio`; the seed then matches its length),
    or ``none``.
    """
    if content not in ("screen", "testsrc"):
        raise ValueError(f"unknown content {content!r}")
    if audio not in ("sine", "speech", "none"):
        raise ValueError(f"unknown audio {audio!r}")
    path.parent.mkdir(parents=True, exist_ok=True)
    seed = path.with_name(f"_seed_{path.stem}.mp4")
    seed_s = gop * 6 / rate
    audio_args: list[str] = []
    speech = path.with_name(f"_speech_{path.stem}.wav")
    if audio == "speech":
        make_speech_audio(speech)
        audio_args = ["-stream_loop", "-1", "-i", str(speech)]
        seed_s = max(seed_s, _probe_duration(speech))
    elif audio == "sine":
        audio_args = ["-f", "lavfi", "-i", f"sine=frequency=220:duration={seed_s}"]
    source = (
        f"testsrc=duration={seed_s}:size={size}:rate=0.2,fps={rate}"
        if content == "screen"
        else f"testsrc=duration={seed_s}:size={size}:rate={rate}"
    )
    cmd = ["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", source]
    cmd += audio_args
    cmd += ["-pix_fmt", "yuv420p", "-c:v", "libx264", "-preset", "veryfast"]
    cmd += ["-crf", "28", "-g", str(gop), "-keyint_min", str(gop)]
    cmd += ["-sc_threshold", "0", "-t", f"{seed_s:.3f}"]
    if audio != "none":
        cmd += ["-c:a", "aac", "-shortest"]
    cmd.append(str(seed))
    subprocess.run(cmd, check=True)
    subprocess.run(
        ["ffmpeg", "-y", "-loglevel", "error", "-stream_loop", "-1", "-i", str(seed)]
        + ["-c", "copy", "-t", str(duration), str(path)],
        check=True,
    )
    seed.unlink()
    if audio == "speech":
        speech.unlink()  # its name would parse as a participant video
    return path


def _probe_duration(path: Path) -> float:
    proc = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration"]
        + ["-of", "csv=p=0", str(path)],
        capture_output=True,
        text=True,
        check=True,
    )
    return float(proc.stdout.strip())


def format_hms(seconds: int) -> str:
    """``H:MM:SS`` (or ``MM:SS`` under an hour), as a sheet cell writes it."""
    h, rest = divmod(int(seconds), 3600)
    m, s = divmod(rest, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


def clip_cells(rows: int, *, duration: int, clip_seconds: int) -> list[str]:
    """*rows* ranges of *clip_seconds*, spread evenly across the whole source.

    Row *i* starts at ``i/rows`` of the usable span, so the last clips sit
    deep in an hours-long file — the seeks short fixtures never exercise.
    """
    clip = max(1, min(int(clip_seconds), int(duration) - 2))
    span = max(0, int(duration) - clip - 2)
    cells = []
    for i in range(rows):
        start = 1 + (span * i) // max(rows, 1)
        cells.append(f"{format_hms(start)}-{format_hms(start + clip)}")
    return cells


def transcripts_section(
    participants: dict[str, int], *, marks_every: int = 0
) -> dict[str, Any]:
    """The ``transcripts`` manifest section with N synthetic segments per id.

    *marks_every* > 0 bookmarks every Nth segment, the shape Composer's
    transcript lane draws.
    """
    source: dict[str, Any] = {}
    for pid, n in participants.items():
        source[pid] = {
            "segments": [
                {
                    "id": f"{pid}:{i}",
                    "start": i * 3.0,
                    "end": i * 3.0 + 2.8,
                    "text": f"Synthetic segment {i} for render benchmarking.",
                }
                for i in range(n)
            ],
            "language": "en",
            "model": "synthetic",
        }
    marks = []
    if marks_every > 0:
        for pid, n in participants.items():
            marks += [
                {
                    "id": f"m_{pid}_{i}",
                    "segment_id": f"{pid}:{i}",
                    "category": "bookmark",
                    "label": None,
                    "severity": None,
                    "created": "2026-01-05T12:00:00+00:00",
                }
                for i in range(0, n, marks_every)
            ]
    return {"source_transcripts": source, "corrections": [], "marks": marks}


def screenspace_section(
    *,
    study: str,
    participant: str,
    video_path: Path,
    events: int,
    duration: float,
    source_size: tuple[int, int] = (1280, 720),
) -> dict[str, Any]:
    """A completed ``change`` task with *events* results and matching events.

    Completed only: an active task opens an SSE stream, which defeats every
    settle heuristic (see ``_ui_fixtures._seed_screenspace``).
    """
    width, height = source_size
    task_id = "bench-task-1"
    step = duration / max(events, 1)
    results = [
        {"timestamp": round(i * step, 3), "magnitude": 0.3 + (i % 7) / 10}
        for i in range(events)
    ]
    event_rows = [
        {
            "id": f"bench-evt-{i}",
            "source_video": video_path.name,
            "participant": participant,
            "detector": "change",
            "event_type": "change",
            "time_in": round(i * step, 3),
            "time_out": round(min(i * step + step * 0.8, duration), 3),
            "confidence": round(0.3 + (i % 7) / 10, 2),
            "metadata": {},
            "excluded": False,
            "task_id": task_id,
            "region": "bench",
        }
        for i in range(events)
    ]
    return {
        "regions": {
            "bench": {
                "x": 0.05,
                "y": 0.05,
                "w": 0.40,
                "h": 0.20,
                "source_width": width,
                "source_height": height,
            }
        },
        "tasks": [
            {
                "id": task_id,
                "type": "change",
                "name": "Change · bench",
                "participant": participant,
                "source_video": video_path.name,
                "video_paths": [str(video_path)],
                "region": "bench",
                "region_coords": {
                    "x": int(width * 0.05),
                    "y": int(height * 0.05),
                    "w": int(width * 0.40),
                    "h": int(height * 0.20),
                },
                "parameters": {},
                "status": "completed",
                "progress": 1.0,
                "priority": 100,
                "result": results,
                "error": None,
                "created_at": "2026-01-05T12:00:00+00:00",
                "completed_at": "2026-01-05T12:01:00+00:00",
            }
        ],
        "events": event_rows,
        "stashes": [],
        "per_participant": {},
        "pins": {},
    }


def write_manifest(out_dir: Path, sections: dict[str, Any]) -> Path:
    """Merge *sections* into ``out_dir/clipgen.json`` by absolute path."""
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "clipgen.json"
    doc: dict[str, Any] = {}
    if path.is_file():
        doc = json.loads(path.read_text(encoding="utf-8"))
    doc.update(sections)
    path.write_text(json.dumps(doc, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def manifest_counts(out_dir: Path) -> dict[str, Any]:
    """What a manifest holds: segments per participant, events, task results."""
    path = out_dir / "clipgen.json"
    if not path.is_file():
        return {"segments": {}, "events": 0, "results": 0}
    doc = json.loads(path.read_text(encoding="utf-8"))
    ts = doc.get("transcripts", {}).get("source_transcripts", {})
    ss = doc.get("screenspace", {})
    return {
        "segments": {pid: len(entry.get("segments", [])) for pid, entry in ts.items()},
        "events": len(ss.get("events", [])),
        "results": sum(len(t.get("result") or []) for t in ss.get("tasks", [])),
    }


def sheet_rows(path: Path) -> int:
    """Observation rows in a sheet written by :func:`write_sheet`."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["Observations"]
    n = sum(1 for row in ws.iter_rows(min_row=6, min_col=5, max_col=5) if row[0].value)
    wb.close()
    return n
